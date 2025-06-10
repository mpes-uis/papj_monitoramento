import pandas as pd
import docx
from docx import Document
import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.base import MIMEBase
from email import encoders
import os
from datetime import datetime
from email.mime.text import MIMEText

# Read the Excel file
df = pd.read_excel('base_acao_exemplo.xlsx')

# Function to replace placeholders in the document
def fill_document(template_path, data, output_path):
    doc = Document(template_path)
    
    # Replace placeholders in the document
    for paragraph in doc.paragraphs:
        for key, value in data.items():
            placeholder = f'{key}'
            if placeholder in paragraph.text:
                paragraph.text = paragraph.text.replace(placeholder, str(value))
    
    doc.save(output_path)

# Function to send email with attachment
def send_email(recipient_email, subject, body, attachment_path):
    # Email configuration
    smtp_server = 'smtp.office365.com'
    smtp_port = 587
    sender_email = ''
    sender_password = ''  # Consider using environment variables for security
    
    # Create message
    msg = MIMEMultipart()
    msg['From'] = sender_email
    msg['To'] = recipient_email
    msg['Subject'] = subject
    
    # Add the email body
    msg.attach(MIMEText(body, 'plain', 'utf-8'))  # Added UTF-8 encoding
    
    # Attach the file
    if os.path.exists(attachment_path):  # Check if file exists
        with open(attachment_path, 'rb') as attachment:
            part = MIMEBase('application', 'vnd.openxmlformats-officedocument.wordprocessingml.document')  # Specific MIME type for .docx
            part.set_payload(attachment.read())
        
        encoders.encode_base64(part)
        
        part.add_header(
            'Content-Disposition',
            f'attachment; filename={os.path.basename(attachment_path)}'  # Removed extra space
        )
        
        msg.attach(part)
    
    # Send email
    try:
        with smtplib.SMTP(smtp_server, smtp_port) as server:  # Using context manager
            server.starttls()
            server.login(sender_email, sender_password)
            server.send_message(msg)  # Using send_message instead of sendmail
        print(f"Email sent successfully to {recipient_email}")
    except Exception as e:
        print(f"Failed to send email to {recipient_email}: {e}")

# Main processing loop
for index, row in df.iterrows():
    # Prepare data dictionary for this row, substituindo valores vazios ou NaN por ""
    data = {
        'id_documento': "" if pd.isna(row['id_documento']) else row['id_documento'],
        'COD_ACAO': "" if pd.isna(row['COD_ACAO']) else row['COD_ACAO'],
        'Responsavel_': "" if pd.isna(row['Responsavel']) else row['Responsavel'],
        'NOME_PJ_CONCATENADO': "" if pd.isna(row['NOME_PJ_CONCATENADO']) else row['NOME_PJ_CONCATENADO'],
        'proced_SEI': "" if pd.isna(row['proced_SEI']) else row['proced_SEI'],
        'TEMA_ind': "" if pd.isna(row['TEMA']) else row['TEMA'],
        'DIRETRIZ_CONSOLIDADA': "" if pd.isna(row['DIRETRIZ_CONSOLIDADA']) else row['DIRETRIZ_CONSOLIDADA'],
        'RESULTADOS_ESPERADOS': "" if pd.isna(row['RESULTADOS_ESPERADOS']) else row['RESULTADOS_ESPERADOS'],
        'Indicador 1: IND_01': "" if pd.isna(row['IND_01']) else row['IND_01'],
        'Indicador 2: IND_02': "" if pd.isna(row['IND_02']) else row['IND_02'],
        'Indicador 3: IND_03': "" if pd.isna(row['IND_03']) else row['IND_03'],
        'Indicador 4: IND_04': "" if pd.isna(row['IND_04']) else row['IND_04'],
        'Indicador 5: IND_05': "" if pd.isna(row['IND_05']) else row['IND_05'],
        'Indicador 6: IND_06': "" if pd.isna(row['IND_06']) else row['IND_06'],
        'Indicador 7: IND_07': "" if pd.isna(row['IND_07']) else row['IND_07'],
        'Indicador 8: IND_08': "" if pd.isna(row['IND_08']) else row['IND_08'],
        'Insira aqui o resultado do indicador 1': "" if pd.isna(row['IND_01']) else 'Insira aqui o resultado do indicador 1',
        'Insira aqui o resultado do indicador 2': "" if pd.isna(row['IND_02']) else 'Insira aqui o resultado do indicador 2',
        'Insira aqui o resultado do indicador 3': "" if pd.isna(row['IND_03']) else 'Insira aqui o resultado do indicador 3',
        'Insira aqui o resultado do indicador 4': "" if pd.isna(row['IND_04']) else 'Insira aqui o resultado do indicador 4',
        'Insira aqui o resultado do indicador 5': "" if pd.isna(row['IND_05']) else 'Insira aqui o resultado do indicador 5',
        'Insira aqui o resultado do indicador 6': "" if pd.isna(row['IND_06']) else 'Insira aqui o resultado do indicador 6',
        'Insira aqui o resultado do indicador 7': "" if pd.isna(row['IND_07']) else 'Insira aqui o resultado do indicador 7',
        'Insira aqui o resultado do indicador 8': "" if pd.isna(row['IND_08']) else 'Insira aqui o resultado do indicador 8'
    }
    # Generate output filename
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    output_filename = f'relatorio_{row["id_documento"]}_{timestamp}.docx'
    
    # Fill document with data
    fill_document('modelo_relatorio.docx', data, output_filename)
    
    # Send email
    email_subject = f'Relatório {row["COD_ACAO"]} - {row["NOME_PJ_CONCATENADO"]}'
    email_body = f'Prezado(a) {row["Responsavel"]},\n\nSegue em anexo o relatório preenchido para a ação {row["COD_ACAO"]}.\n\nAtenciosamente,\nSistema Automático'
    
    send_email(row['E-mail'], email_subject, email_body, output_filename)

print("All reports generated and emails sent.")
